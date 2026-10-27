from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

import sys, os
VERSION = "1.0.0"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BLANK = len(sys.argv) > 1 and sys.argv[1] == "blank"
N1, N2 = ("Personne 1", "Personne 2") if not BLANK else ("Personne 1", "Personne 2")
wb = Workbook()
F = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
INP = PatternFill("solid", fgColor="FFF9E5")
TOT = PatternFill("solid", fgColor="D9E2F3")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 "€";-#,##0.00 "€";-'
PCT = '0%'

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(name=F, size=14, bold=True, color="1F3864")
    ws["A2"] = sub; ws["A2"].font = Font(name=F, size=9, italic=True, color="595959")

def header(ws, row, labels):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = 42

def style(c, inp=False, fmt=None, bold=False):
    c.font = Font(name=F, size=10, bold=bold, color="0000FF" if inp else "000000")
    c.border = BOX
    if inp: c.fill = INP
    if fmt: c.number_format = fmt

def widths(ws, ws_w):
    for k, v in ws_w.items(): ws.column_dimensions[k].width = v

P = "'Paramètres'!"
# ---------- Paramètres
wp = wb.active; wp.title = "Paramètres"
title(wp, "Paramètres du foyer", "Cellules jaunes à police bleue = à modifier. Tout le reste se calcule.")
for r, (l, v) in enumerate([("Personne 1", "Personne 1" if BLANK else "Alex"), ("Personne 2", "Personne 2" if BLANK else "Sam"),
        ("Compte perso de la Personne 1", None if BLANK else "Compte Alex"), ("Compte perso de la Personne 2", None if BLANK else "Compte Sam")], 4):
    wp.cell(r, 1, l).font = Font(name=F, size=10, bold=True)
    style(wp.cell(r, 2, v), inp=True)
wp["A9"] = "Part de la Personne 1 dans les salaires (pour « Prorata salaires »)"
wp["A9"].font = Font(name=F, size=10, bold=True); wp["A9"].alignment = Alignment(wrap_text=True)
wp["B9"] = ('=IFERROR(SUMIFS(Revenus!$F$5:$F$24,Revenus!$B$5:$B$24,"Salaire",Revenus!$C$5:$C$24,$B$4)'
            '/SUMIFS(Revenus!$F$5:$F$24,Revenus!$B$5:$B$24,"Salaire"),0.5)')
style(wp["B9"], fmt="0.0%")
wp["A10"] = "Hypothèse : 50 % si aucun salaire n'est saisi."; wp["A10"].font = Font(name=F, size=9, italic=True, color="595959")
wp["D3"] = "Périodicité"; wp["E3"] = "Nb de mois"
for c in (wp["D3"], wp["E3"]): c.font = Font(name=F, bold=True, size=10)
for r, (p, n) in enumerate([("Mensuelle", 1), ("Trimestrielle", 3), ("Semestrielle", 6), ("Annuelle", 12)], 4):
    style(wp.cell(r, 4, p)); style(wp.cell(r, 5, n))
wp["A12"] = "Modes de répartition possibles : 50/50 · Prorata salaires · Personne 1 seul · Personne 2 seul"
wp["A12"].font = Font(name=F, size=9, italic=True, color="595959")
wp["A14"] = "Budget personnel (le « reste libre »)"; wp["A14"].font = Font(name=F, bold=True, size=11, color="1F3864")
wp["B14"] = "=$B$4"; wp["C14"] = "=$B$5"
for c in (wp["B14"], wp["C14"]): c.font = Font(name=F, bold=True, size=10); c.alignment = Alignment(horizontal="center")
rows = [("Charges perso : mode (Détail = onglet « Charges perso » ; Forfait = somme arbitraire)", "Détail", "Forfait", None),
        ("Forfait charges perso (€/mois), utilisé si mode Forfait", 0, 350, EUR),
        ("Épargne perso : mode (Pourcentage ou Montant)", "Pourcentage", "Montant", None),
        ("Épargne perso : pourcentage du reste (si Pourcentage)", 0.3, 0, "0%"),
        ("Épargne perso : montant €/mois (si Montant)", 0, 200, EUR)]
if BLANK:
    rows = [(l, a if isinstance(a, str) else 0, b if isinstance(b, str) else 0, f) for l, a, b, f in rows]
    rows[0] = (rows[0][0], "Détail", "Détail", None); rows[2] = (rows[2][0], "Pourcentage", "Pourcentage", None)
for r, (l, v1, v2, fmt) in enumerate(rows, 15):
    wp.cell(r, 1, l).font = Font(name=F, size=10); wp.cell(r, 1).alignment = Alignment(wrap_text=True); wp.row_dimensions[r].height = 28
    style(wp.cell(r, 2, v1), inp=True, fmt=fmt); style(wp.cell(r, 3, v2), inp=True, fmt=fmt)
for rng, f in [("B15:C15", '"Détail,Forfait"'), ("B17:C17", '"Pourcentage,Montant"')]:
    v = DataValidation(type="list", formula1=f, allow_blank=False); wp.add_data_validation(v); v.add(rng)
wp["A20"] = "Si l'un de vous ne veut pas détailler ses comptes : mode Forfait + une somme. Ses lignes éventuelles dans « Charges perso » sont alors ignorées."
wp["A20"].font = Font(name=F, size=9, italic=True, color="595959")
wp["A22"] = "Provision imprévus (% des salaires du foyer)"; wp["A22"].font = Font(name=F, bold=True, size=10)
style(wp["B22"], inp=True, fmt="0.0%"); wp["B22"] = 0 if BLANK else 0.03
wp["A23"] = "Alimente la ligne « Provision imprévus » de l'onglet Charges (réparties 50/50, sur le compte joint indiqué)."
wp["A23"].font = Font(name=F, size=9, italic=True, color="595959")
wp["A25"] = "Version du modèle"; wp["A25"].font = Font(name=F, size=10, bold=True)
wp["B25"] = VERSION; style(wp["B25"])
wp["A26"] = "Voir MIGRATION.md pour migrer un fichier d'une version à l'autre."
wp["A26"].font = Font(name=F, size=9, italic=True, color="595959")
widths(wp, {"A": 52, "B": 22, "C": 22, "D": 16, "E": 12})
wp.row_dimensions[9].height = 30

# ---------- Comptes
wc = wb.create_sheet("Comptes")
title(wc, "Comptes bancaires", "Paramétrez ici vos comptes (10 lignes). Le nom saisi ici alimente les listes déroulantes des autres onglets.")
header(wc, 4, ["Nom du compte", "Type", "Titulaire(s)", "Usage"])
comptes = [("Compte Alex", "Personnel", "Alex", "Salaire d'Alex, dépenses perso"),
           ("Compte Sam", "Personnel", "Sam", "Salaire de Sam, dépenses perso"),
           ("Joint Logement", "Joint", "Alex & Sam", "Emprunt, assurances, énergie, taxe foncière"),
           ("Joint Vie courante", "Joint", "Alex & Sam", "Courses, voiture, enfants, aides CAF"),
           ("Joint Abonnements", "Joint", "Alex & Sam", "Internet, streaming, etc."),
           ("Livret A", "Épargne", "Alex & Sam", "Épargne (reçoit des virements)"),
           ("Livret imprévus", "Épargne", "Alex & Sam", "Provision pour les dépenses imprévues"),
           ("Livret enfant 1", "Épargne", "Enfant 1", "Épargne de l'enfant 1"),
           ("Livret enfant 2", "Épargne", "Enfant 2", "Épargne de l'enfant 2")]
if BLANK: comptes = []
for r in range(5, 15):
    for c in range(1, 5):
        v = comptes[r-5][c-1] if r-5 < len(comptes) else None
        style(wc.cell(r, c, v), inp=True)
dv = DataValidation(type="list", formula1='"Joint,Personnel,Épargne"', allow_blank=True); wc.add_data_validation(dv); dv.add("B5:B14")
widths(wc, {"A": 26, "B": 14, "C": 18, "D": 48})
wc.freeze_panes = "A5"

# ---------- Charges
wch = wb.create_sheet("Charges")
title(wch, "Charges (dépenses récurrentes)",
      "Remplir les colonnes jaunes (50 lignes). Montant = montant par période ; le mensuel et les parts se calculent. Virement : indiquer le compte destinataire (liste libre). Dépenses irrégulières : saisir une enveloppe annuelle, elle est lissée par mois.")
header(wch, 4, ["Libellé", "Catégorie", "Nature", "Périodicité", "Montant (€ par période)", "Montant mensuel",
                "Mode de paiement", "Compte débité", "Compte destinataire (si virement)", "Jour du mois", "Répartition",
                "Part P1 (%)", None, None])
wch["M4"] = '="Part "&' + P + '$B$4&" (€/mois)"'; wch["N4"] = '="Part "&' + P + '$B$5&" (€/mois)"'
ch = [("Emprunt immobilier", "Logement", "Fixe", "Mensuelle", 1150, "Prélèvement", "Joint Logement", None, 5, "Prorata salaires"),
      ("Assurance emprunt", "Logement", "Fixe", "Mensuelle", 45, "Prélèvement", "Joint Logement", None, 5, "50/50"),
      ("Taxe foncière", "Logement", "Fixe", "Annuelle", 1300, "Prélèvement", "Joint Logement", None, 15, "Prorata salaires"),
      ("Électricité maison", "Énergie", "Variable", "Mensuelle", 130, "Prélèvement", "Joint Logement", None, 12, "50/50"),
      ("Électricité garage / local", "Énergie", "Variable", "Mensuelle", 40, "Prélèvement", "Joint Abonnements", None, 14, "50/50"),
      ("Internet / box", "Abonnements", "Fixe", "Mensuelle", 30, "Prélèvement", "Joint Abonnements", None, 10, "50/50"),
      ("Assurance auto (voiture d'Alex)", "Transport", "Fixe", "Trimestrielle", 210, "Prélèvement", "Joint Vie courante", None, 1, "Personne 1 seul"),
      ("Forfait mobile de Sam", "Abonnements", "Fixe", "Mensuelle", 28, "Prélèvement", "Compte Sam", None, 8, "Personne 2 seul"),
      ("Courses alimentaires", "Vie courante", "Variable", "Mensuelle", 600, "Carte", "Joint Vie courante", None, None, "50/50"),
      ("Épargne mensuelle", "Épargne", "Fixe", "Mensuelle", 150, "Virement", "Joint Vie courante", "Livret A", 3, "50/50"),
      ("Cantine des enfants", "Enfants", "Variable", "Mensuelle", 90, "Virement", "Joint Vie courante", "Mairie (cantine)", 28, "50/50"),
      ("Vêtements enfants (enveloppe lissée)", "Enfants", "Variable", "Annuelle", 600, "Carte", "Joint Vie courante", None, None, "50/50"),
      ("Médecin / pharmacie (provision)", "Santé", "Variable", "Mensuelle", 30, "Chèque", "Joint Vie courante", None, None, "50/50"),
      ("Cadeaux & fêtes (enveloppe lissée)", "Loisirs", "Variable", "Annuelle", 500, "Carte", "Joint Vie courante", None, None, "50/50"),
      ("Provision imprévus (% des salaires)", "Imprévus", "Variable", "Mensuelle", "=PROV", "Virement", "Joint Vie courante", "Livret imprévus", 3, "50/50"),
      ("Épargne enfant 1", "Épargne", "Fixe", "Mensuelle", 50, "Virement", "Joint Vie courante", "Livret enfant 1", 4, "50/50"),
      ("Épargne enfant 2", "Épargne", "Fixe", "Mensuelle", 50, "Virement", "Joint Vie courante", "Livret enfant 2", 4, "50/50"),
      ("Aide à domicile (retrait espèces)", "Services", "Fixe", "Mensuelle", 120, "Espèces", "Joint Vie courante", None, 1, "50/50"),
      ("Argent de poche enfants (retrait espèces)", "Enfants", "Fixe", "Mensuelle", 30, "Espèces", "Joint Vie courante", None, 1, "50/50"),
      ("Marché & petits commerces (retrait espèces)", "Vie courante", "Variable", "Mensuelle", 60, "Espèces", "Joint Vie courante", None, None, "50/50")]
if BLANK: ch = []
for r in range(5, 55):
    d = ch[r-5] if r-5 < len(ch) else None
    vals = [d[0], d[1], d[2], d[3], d[4], None, d[5], d[6], d[7], d[8], d[9]] if d else [None]*11
    for c, v in enumerate(vals, 1):
        if c == 6: continue
        style(wch.cell(r, c, v), inp=True, fmt=EUR if c == 5 else None)
    wch.cell(r, 6, f'=IF(A{r}="","",IFERROR(E{r}/INDEX({P}$E$4:$E$7,MATCH(D{r},{P}$D$4:$D$7,0)),0))')
    wch.cell(r, 12, f'=IF(A{r}="","",IF(K{r}="50/50",0.5,IF(K{r}="Prorata salaires",{P}$B$9,IF(K{r}="Personne 1 seul",1,IF(K{r}="Personne 2 seul",0,0.5)))))')
    wch.cell(r, 13, f'=IF(A{r}="","",F{r}*L{r})'); wch.cell(r, 14, f'=IF(A{r}="","",F{r}-M{r})')
    style(wch.cell(r, 6), fmt=EUR); style(wch.cell(r, 12), fmt=PCT)
    style(wch.cell(r, 13), fmt=EUR); style(wch.cell(r, 14), fmt=EUR)
for r in range(5, 55):
    if wch.cell(r, 5).value == "=PROV":
        wch.cell(r, 5).value = '=' + P + '$B$22*SUMIFS(Revenus!$F$5:$F$24,Revenus!$B$5:$B$24,"Salaire")'
        wch.cell(r, 5).font = Font(name=F, size=10, color="000000")
wch["A55"] = "TOTAL mensuel"; wch["A55"].font = Font(name=F, bold=True)
for col in "FMN":
    wch[f"{col}55"] = f"=SUM({col}5:{col}54)"; style(wch[f"{col}55"], fmt=EUR, bold=True); wch[f"{col}55"].fill = TOT
for rng, f in [("C5:C54", '"Fixe,Variable"'), ("D5:D54", '"Mensuelle,Trimestrielle,Semestrielle,Annuelle"'),
               ("G5:G54", '"Prélèvement,Virement,Carte,Chèque,Espèces"'), ("H5:H54", "Comptes!$A$5:$A$14"),
               ("K5:K54", '"50/50,Prorata salaires,Personne 1 seul,Personne 2 seul"')]:
    v = DataValidation(type="list", formula1=f, allow_blank=True); wch.add_data_validation(v); v.add(rng)
widths(wch, {"A": 32, "B": 14, "C": 11, "D": 14, "E": 14, "F": 13, "G": 14, "H": 20, "I": 22, "J": 9, "K": 18, "L": 9, "M": 14, "N": 14})
wch.freeze_panes = "B5"

# ---------- Revenus
wr = wb.create_sheet("Revenus")
title(wr, "Revenus et aides", "Salaires (arrivent sur le compte de chacun) et aides type CAF (peuvent arriver sur un compte joint). 20 lignes.")
header(wr, 4, ["Libellé", "Type", "Personne", "Montant (€ par période)", "Périodicité", "Montant mensuel", "Compte de réception", "Jour du mois"])
rv = [("Salaire Alex", "Salaire", "Alex", 2400, "Mensuelle", "Compte Alex", 28),
      ("Salaire Sam", "Salaire", "Sam", 1900, "Mensuelle", "Compte Sam", 30),
      ("Allocations familiales (CAF)", "Aide", "Commun", 140, "Mensuelle", "Joint Vie courante", 5),
      ("Prime d'activité (CAF)", "Aide", "Commun", 270, "Trimestrielle", "Joint Vie courante", 5)]
if BLANK: rv = []
for r in range(5, 25):
    d = rv[r-5] if r-5 < len(rv) else None
    vals = [d[0], d[1], d[2], d[3], d[4], None, d[5], d[6]] if d else [None]*8
    for c, v in enumerate(vals, 1):
        if c == 6: continue
        style(wr.cell(r, c, v), inp=True, fmt=EUR if c == 4 else None)
    wr.cell(r, 6, f'=IF(A{r}="","",IFERROR(D{r}/INDEX({P}$E$4:$E$7,MATCH(E{r},{P}$D$4:$D$7,0)),0))'); style(wr.cell(r, 6), fmt=EUR)
wr["A25"] = "TOTAL mensuel"; wr["A25"].font = Font(name=F, bold=True)
wr["F25"] = "=SUM(F5:F24)"; style(wr["F25"], fmt=EUR, bold=True); wr["F25"].fill = TOT
for rng, f in [("B5:B24", '"Salaire,Aide,Autre"'), ("C5:C24", f'"{{}}"'), ("E5:E24", '"Mensuelle,Trimestrielle,Semestrielle,Annuelle"'),
               ("G5:G24", "Comptes!$A$5:$A$14")]:
    if rng.startswith("C"): continue
    v = DataValidation(type="list", formula1=f, allow_blank=True); wr.add_data_validation(v); v.add(rng)
widths(wr, {"A": 32, "B": 10, "C": 12, "D": 14, "E": 14, "F": 13, "G": 22, "H": 9})
wr.freeze_panes = "B5"

# ---------- Charges perso
wcp = wb.create_sheet("Charges perso")
title(wcp, "Charges personnelles (facultatif)", "Dépenses propres à chacun. Utilisé seulement si le mode « Détail » est choisi pour la personne dans Paramètres.")
header(wcp, 4, ["Libellé", "Personne", "Catégorie", "Périodicité", "Montant (€ par période)", "Montant mensuel", "Mode de paiement", "Compte débité", "Jour du mois"])
cp_ex = [("Salle de sport", "Alex", "Loisirs", "Mensuelle", 35, "Prélèvement", "Compte Alex", 10),
         ("Coiffeur", "Alex", "Soins", "Trimestrielle", 120, "Carte", "Compte Alex", None),
         ("Abonnement musique", "Alex", "Abonnements", "Mensuelle", 11, "Prélèvement", "Compte Alex", 15)]
if BLANK: cp_ex = []
for r in range(5, 45):
    d = cp_ex[r-5] if r-5 < len(cp_ex) else (None,)*8
    for c, v in zip((1, 2, 3, 4, 5, 7, 8, 9), d):
        style(wcp.cell(r, c, v), inp=True, fmt=EUR if c == 5 else None)
    wcp.cell(r, 6, f'=IF(A{r}="","",IFERROR(E{r}/INDEX({P}$E$4:$E$7,MATCH(D{r},{P}$D$4:$D$7,0)),0))'); style(wcp.cell(r, 6), fmt=EUR)
wcp["A45"] = "TOTAL mensuel"; wcp["A45"].font = Font(name=F, bold=True)
wcp["F45"] = "=SUM(F5:F44)"; style(wcp["F45"], fmt=EUR, bold=True); wcp["F45"].fill = TOT
for rng, f in [("G5:G44", '"Prélèvement,Carte,Chèque,Espèces,Virement"'), ("H5:H44", "Comptes!$A$5:$A$14"), ("B5:B44", P + "$B$4:$B$5"), ("D5:D44", '"Mensuelle,Trimestrielle,Semestrielle,Annuelle"')]:
    v = DataValidation(type="list", formula1=f, allow_blank=True); wcp.add_data_validation(v); v.add(rng)
widths(wcp, {"A": 32, "B": 12, "C": 14, "D": 14, "E": 14, "F": 13, "G": 16, "H": 20, "I": 9})
wcp.freeze_panes = "A5"

# ---------- Imprévus
from datetime import date
wi = wb.create_sheet("Imprévus")
title(wi, "Suivi des imprévus", "Notez chaque dépense imprévue ici (date, libellé, montant). Le solde de la provision se met à jour tout seul.")
lab = [("Provision mensuelle (ligne « Imprévus » de Charges)", '=SUMIFS(Charges!$F$5:$F$54,Charges!$B$5:$B$54,"Imprévus")', EUR, False),
       ("Début du suivi (1er du mois)", None if BLANK else date(2026, 9, 1), "dd/mm/yyyy", True),
       ("Solde de départ du livret imprévus", 0, EUR, True),
       ("Mois provisionnés (mois en cours inclus)", '=IF(B5="",0,MAX(0,(YEAR(TODAY())-YEAR(B5))*12+MONTH(TODAY())-MONTH(B5)+1))', "0", False),
       ("Total provisionné", "=B6+B4*B7", EUR, False),
       ("Total dépensé en imprévus", "=SUM(C13:C62)", EUR, False),
       ("Solde de la provision", "=B8-B9", EUR, False)]
for i, (l, v, fmt, inp) in enumerate(lab, 4):
    wi.cell(i, 1, l).font = Font(name=F, size=10, bold=(i == 10))
    style(wi.cell(i, 2, v), inp=inp, fmt=fmt, bold=(i == 10))
wi["B10"].fill = TOT
wi["A10"].font = Font(name=F, size=10, bold=True)
wi["A11"] = "Solde négatif = la provision n'a pas suffi : à combler avec l'épargne ou le reste libre."
wi["A11"].font = Font(name=F, size=9, italic=True, color="595959")

header(wi, 12, ["Date", "Libellé", "Montant (€)", "Payé depuis"])
ex = [(date(2026, 9, 20), "Pare-brise", 220, "Livret imprévus"), (date(2026, 10, 5), "Dentiste", 85, "Livret imprévus")]
if BLANK: ex = []
for r in range(13, 63):
    d = ex[r-13] if r-13 < len(ex) else (None,)*4
    for c, v in enumerate(d, 1):
        style(wi.cell(r, c, v), inp=True, fmt="dd/mm/yyyy" if c == 1 else (EUR if c == 3 else None))
v = DataValidation(type="list", formula1="Comptes!$A$5:$A$14", allow_blank=True); wi.add_data_validation(v); v.add("D13:D62")
widths(wi, {"A": 48, "B": 28, "C": 14, "D": 22})
wi.freeze_panes = "A13"

# ---------- Synthèse
ws = wb.create_sheet("Synthèse", 0)
title(ws, "Synthèse mensuelle", "Tout est calculé à partir des onglets Charges, Revenus et Comptes. Ne rien saisir ici.")
ws["A3"] = "1. Par personne"; ws["A3"].font = Font(name=F, bold=True, size=11, color="1F3864")
header(ws, 4, ["Personne", "Salaires mensuels", "Participations aux comptes joints", "Charges payées directement (compte perso)",
               "Total charges supportées", "Taux d'effort", "Reste après charges communes",
               "Charges perso (détail ou forfait)", "Reste après charges perso", "Épargne perso", "Reste libre (loisirs…)"])
for r, nm, cp, col, pc in [(5, "$B$4", "$B$6", "E", "B"), (6, "$B$5", "$B$7", "F", "C")]:
    ws.cell(r, 1, f"={P}{nm}")
    ws.cell(r, 2, f'=SUMIFS(Revenus!$F$5:$F$24,Revenus!$B$5:$B$24,"Salaire",Revenus!$C$5:$C$24,A{r})')
    ws.cell(r, 3, f"=SUM({col}$13:{col}$22)")
    ws.cell(r, 4, f'=IF({P}{cp}="",0,IFERROR(INDEX($C$13:$C$22,MATCH({P}{cp},$A$13:$A$22,0)),0))')
    ws.cell(r, 5, f"=C{r}+D{r}"); ws.cell(r, 6, f"=IF(B{r}=0,0,E{r}/B{r})")
    ws.cell(r, 7, f'=IF({P}{cp}="",0,IFERROR(INDEX($J$13:$J$22,MATCH({P}{cp},$A$13:$A$22,0)),0))')
    ws.cell(r, 8, f"=IF({P}{pc}$15=\"Forfait\",{P}{pc}$16,SUMIFS('Charges perso'!$F$5:$F$44,'Charges perso'!$B$5:$B$44,A{r}))")
    ws.cell(r, 9, f"=G{r}-H{r}")
    ws.cell(r, 10, f"=MAX(0,IF({P}{pc}$17=\"Pourcentage\",{P}{pc}$18*MAX(0,I{r}),{P}{pc}$19))")
    ws.cell(r, 11, f"=I{r}-J{r}")
ws["A7"] = "Total foyer"
for col in "BCDEGHIJK": ws[f"{col}7"] = f"=SUM({col}5:{col}6)"
ws["F7"] = "=IF(B7=0,0,E7/B7)"
for r in (5, 6, 7):
    for c in range(1, 12):
        style(ws.cell(r, c), fmt=PCT if c == 6 else (None if c == 1 else EUR), bold=(r == 7))
        if r == 7: ws.cell(r, c).fill = TOT
ws["A9"] = "2. Par compte : ce que chacun doit virer chaque mois"; ws["A9"].font = Font(name=F, bold=True, size=11, color="1F3864")
ws["A10"] = ("Un compte joint est alimenté par les virements des deux personnes (colonnes E et F) = charges du compte − aides reçues dessus, "
             "réparties selon la répartition de chaque charge.")
ws["A10"].font = Font(name=F, size=9, italic=True, color="595959")
header(ws, 12, ["Compte", "Type", "Charges débitées / virements émis", "Aides & autres recettes reçues", None, None,
                "Salaires reçus", "Virements reçus (d'autres comptes)", "Participations versées (vers comptes joints)",
                "Solde mensuel prévisionnel"])
ws["E12"] = '="À virer par "&' + P + '$B$4'; ws["F12"] = '="À virer par "&' + P + '$B$5'
for i in range(10):
    r = 13 + i; cr = 5 + i
    ws.cell(r, 1, f'=IF(Comptes!A{cr}="","",Comptes!A{cr})')
    ws.cell(r, 2, f'=IF($A{r}="","",Comptes!B{cr})')
    ws.cell(r, 3, f'=IF($A{r}="","",SUMIFS(Charges!$F$5:$F$54,Charges!$H$5:$H$54,$A{r}))')
    ws.cell(r, 4, f'=IF($A{r}="","",SUMIFS(Revenus!$F$5:$F$24,Revenus!$G$5:$G$24,$A{r},Revenus!$B$5:$B$24,"<>Salaire"))')
    for col, src in (("E", "M"), ("F", "N")):
        ws[f"{col}{r}"] = (f'=IF($A{r}="","",IF($B{r}="Joint",IF($C{r}=0,0,($C{r}-$D{r})*SUMIFS(Charges!${src}$5:${src}$54,'
                           f'Charges!$H$5:$H$54,$A{r})/$C{r}),0))')
    ws.cell(r, 7, f'=IF($A{r}="","",SUMIFS(Revenus!$F$5:$F$24,Revenus!$G$5:$G$24,$A{r},Revenus!$B$5:$B$24,"Salaire"))')
    ws.cell(r, 8, f'=IF($A{r}="","",SUMIFS(Charges!$F$5:$F$54,Charges!$I$5:$I$54,$A{r},Charges!$G$5:$G$54,"Virement"))')
    ws.cell(r, 9, f'=IF($A{r}="","",IF($A{r}={P}$B$6,SUM($E$13:$E$22),IF($A{r}={P}$B$7,SUM($F$13:$F$22),0)))')
    ws.cell(r, 10, f'=IF($A{r}="","",G{r}+D{r}+H{r}+IF(B{r}="Joint",E{r}+F{r},0)-C{r}-I{r})')
    for c in range(1, 11): style(ws.cell(r, c), fmt=None if c <= 2 else EUR)
    for c in (5, 6): ws.cell(r, c).font = Font(name=F, size=10, bold=True)
ws["A23"] = "Total"
for col in "CDEFGHIJ": ws[f"{col}23"] = f"=SUM({col}13:{col}22)"
for c in range(1, 11):
    style(ws.cell(23, c), fmt=None if c <= 2 else EUR, bold=True); ws.cell(23, c).fill = TOT
ws["A27"] = "3. Charges communes : fixe vs variable (mensuel)"; ws["A27"].font = Font(name=F, bold=True, size=11, color="1F3864")
for i, (l, k) in enumerate([("Fixe", "Fixe"), ("Variable / estimé", "Variable")]):
    ws.cell(28+i, 1, l); ws.cell(28+i, 2, f'=SUMIFS(Charges!$F$5:$F$54,Charges!$C$5:$C$54,"{k}")')
    ws.cell(28+i, 3, f"=IF($B$30=0,0,B{28+i}/$B$30)")
    style(ws.cell(28+i, 1)); style(ws.cell(28+i, 2), fmt=EUR); style(ws.cell(28+i, 3), fmt=PCT)
ws["A30"] = "Total"; ws["B30"] = "=B28+B29"; ws["C30"] = "=C28+C29"
style(ws["A30"], bold=True); style(ws["B30"], fmt=EUR, bold=True); style(ws["C30"], fmt=PCT, bold=True)
for c in "ABC": ws[f"{c}30"].fill = TOT
ws["A25"] = ("Lecture : un compte joint à solde 0 est équilibré. Un compte perso à solde positif = reste à vivre ; négatif = salaire insuffisant. "
             "Les aides sont réparties sur les deux personnes au prorata de leurs parts de charges (hypothèse).")
ws["A25"].font = Font(name=F, size=9, italic=True, color="595959")
widths(ws, {"A": 24, "B": 14, "C": 18, "D": 20, "E": 16, "F": 16, "G": 18, "H": 18, "I": 20, "J": 18, "K": 18})
ws.row_dimensions[4].height = 55; ws.row_dimensions[12].height = 55

out = os.path.join(ROOT, "suivi_comptes_vierge.xlsx" if BLANK else "suivi_comptes_example.xlsx")
wb.save(out); print(out)
