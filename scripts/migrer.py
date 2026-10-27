"""Migre un fichier réel vers le modèle courant : copie les zones de saisie (voir MIGRATION.md).

Usage : python3 scripts/migrer.py ancien.xlsx nouveau.xlsx
Le fichier de sortie contient des données réelles : ne JAMAIS le placer dans le dépôt Git.
"""
import sys, os
from openpyxl import load_workbook

# Zones de saisie du modèle v1.0.0 : feuille -> plages recopiées (les colonnes calculées sont exclues).
ZONES = {
    "Paramètres": ["B4:B7", "B15:C19", "B22:B22"],
    "Comptes": ["A5:D14"],
    "Charges": ["A5:E54", "G5:K54"],
    "Revenus": ["A5:E24", "G5:H24"],
    "Charges perso": ["A5:E44", "G5:I44"],
    "Imprévus": ["B5:B6", "A13:D62"],
}
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MODELE = os.path.join(ROOT, "suivi_comptes_vierge.xlsx")

def main(ancien, nouveau):
    src, dst = load_workbook(ancien), load_workbook(MODELE)
    print("Version source :", src["Paramètres"]["B25"].value or "inconnue (avant 1.0.0)")
    for feuille, plages in ZONES.items():
        for plage in plages:
            for ligne in src[feuille][plage]:
                for c in ligne:
                    dst[feuille][c.coordinate].value = c.value
    dst.save(nouveau)
    print("Écrit :", nouveau, "- recalculer (Excel/Sheets) puis vérifier la Synthèse.")

if __name__ == "__main__":
    main(*sys.argv[1:3])
